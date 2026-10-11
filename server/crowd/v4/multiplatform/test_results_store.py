"""Offline business result contract tests. No model or source requests."""
import copy
import csv
import io
import json
from pathlib import Path
import sqlite3
import tempfile
import threading
import unittest

from results_store import ResultsStore


def imported(run='run-1', body='Original text', likes=1):
    return {'run_id': run, 'source_kind': 'mediacrawler_api', 'coverage': 'unverified',
            'source': {'adapter_version': 'test'}, 'records': [
                {'kind': 'content', 'platform': 'xiaohongshu', 'content_id': 'n1',
                 'author_id': 'a1', 'title': 'Title', 'body': body,
                 'published_at_reported': 1704067200000,
                 'metrics': {'liked_count': {'reported_value': likes}},
                 'origin': {'run_id': run, 'line': 1}},
                {'kind': 'comment', 'platform': 'xiaohongshu', 'content_id': 'n1',
                 'comment_id': 'c1', 'parent_comment_id': '0', 'text': 'Root'},
                {'kind': 'comment', 'platform': 'xiaohongshu', 'content_id': 'n1',
                 'comment_id': 'c2', 'parent_comment_id': 'c1', 'text': 'Reply'},
                {'kind': 'author', 'platform': 'xiaohongshu', 'author_id': 'a1',
                 'display_name': 'Author'}],
            'media': [{'content_id': 'n1', 'path': 'output/xhs/media/n1.jpg',
                       'bytes': 10, 'status': 'downloaded_by_external_executor'}]}


class ResultsTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.store = ResultsStore(self.temporary.name)

    def test_idempotent_import_and_restart(self):
        first = self.store.ingest(imported())
        self.assertEqual(first['added_versions'], 5)
        again = ResultsStore(self.temporary.name).ingest(imported())
        self.assertTrue(again['duplicate_import'])
        self.assertEqual(again['results_revision'], first['results_revision'])
        self.assertEqual(self.store.path.stat().st_mode & 0o777, 0o600)

    def test_zhihu_reported_types_keep_distinct_entities_and_reject_ambiguous_evaluation(self):
        records=[{'kind':'content','platform':'zhihu','content_id':123,'content_type':kind,'body':kind,
                  'record_hash':kind+'-unchanged'} for kind in ('answer','article','zvideo')]
        self.store.ingest({'records':records})
        detail=self.store.detail('zhihu','123')
        self.assertEqual(detail['stats']['contents'],3)
        self.assertEqual(detail['content_identity_status'],'ambiguous_types')
        self.assertEqual(detail['observed_content_types'],['answer','article','zvideo'])
        self.assertEqual({r['entity_ref'] for r in detail['contents']},{'zhihu:content:123:'+kind for kind in ('answer','article','zvideo')})
        with self.assertRaisesRegex(ValueError,'content_evaluation_requires_one_content'):
            self.store.evaluate('zhihu','123',4)
        unknown={'kind':'comment','platform':'zhihu','content_id':'123','comment_id':'c','text':'unknown parent scope'}
        self.store.ingest({'records':[unknown]})
        comment=self.store.detail('zhihu','123')['comments'][0]
        self.assertNotIn('content_type',comment)

    def test_exact_entity_detail_versions_evaluation_and_unknown_relationships(self):
        answer={'kind':'content','platform':'zhihu','content_id':'123','content_type':'answer','body':'answer v1','record_hash':'original-answer-hash'}
        article={**answer,'content_type':'article','body':'article','record_hash':'original-article-hash'}
        comment={'kind':'comment','platform':'zhihu','content_id':'123','comment_id':'unknown','text':'unknown comment'}
        comments=[comment,{**comment,'comment_id':'answer-c','content_type':'answer'},
                  {**comment,'comment_id':'article-c','platform_fields':{'content_type':'article'}},
                  {**comment,'comment_id':'conflict','source_content_type':'article','content_type':'answer'}]
        media=[{'kind':'media','platform':'zhihu','content_id':'123','media_id':kind,'path':kind+'.jpg',
                'source_content_type':kind if kind!='unknown' else None} for kind in ('answer','article','unknown')]
        first=self.store.ingest({'run_id':'typed-first','records':[answer,article,*comments,*media]})['results_revision']
        self.store.ingest({'run_id':'typed-next','records':[{**answer,'body':'answer v2'}]})
        with sqlite3.connect(str(self.store.path)) as db:
            before={table:list(db.execute('SELECT * FROM '+table)) for table in ('versions','observations','imports')}
        ref='zhihu:content:123:answer';detail=self.store.detail('zhihu','123',entity_ref=ref)
        self.assertEqual(detail['selected_entity_ref'],ref);self.assertEqual([c['body'] for c in detail['contents']],['answer v2'])
        self.assertEqual([v['record']['body'] for v in detail['versions']],['answer v1','answer v2'])
        self.assertEqual(len(detail['observations']),2)
        self.assertEqual([c['comment_id'] for c in detail['comments']],['answer-c'])
        self.assertEqual({c['comment_id'] for c in detail['unassigned_comments']},{'unknown','conflict'})
        self.assertEqual([m['media_id'] for m in detail['media']],['answer'])
        self.assertEqual([m['media_id'] for m in detail['unassigned_media']],['unknown'])
        article_detail=self.store.detail('zhihu','123',entity_ref='zhihu:content:123:article')
        self.assertEqual([c['comment_id'] for c in article_detail['comments']],['article-c'])
        self.assertEqual({c['comment_id'] for c in article_detail['unassigned_comments']},{'unknown','conflict'})
        self.assertEqual(len(self.store.detail('zhihu','123')['comments']),4)
        evaluation=self.store.evaluate('zhihu','123',5,['selected'],entity_ref=ref)
        self.assertEqual(evaluation['entity_ref'],ref)
        self.assertEqual(self.store.detail('zhihu','123',entity_ref=ref)['contents'][0]['owner_evaluations'][0]['value'],5)
        self.assertEqual(self.store.detail('zhihu','123',entity_ref='zhihu:content:123:article')['contents'][0]['owner_evaluations'],[])
        self.assertEqual(self.store.detail('zhihu','123',first,ref)['contents'][0]['body'],'answer v1')
        with sqlite3.connect(str(self.store.path)) as db:
            for table,rows in before.items():self.assertEqual(list(db.execute('SELECT * FROM '+table)),rows)
        with self.assertRaisesRegex(ValueError,'content_evaluation_requires_one_content'):self.store.evaluate('zhihu','123',4)

    def test_entity_selector_boundaries_and_non_zhihu_compatibility(self):
        revision=self.store.ingest(imported())['results_revision'];ref='xiaohongshu:content:n1'
        self.assertEqual(self.store.detail('xhs','n1',entity_ref=ref)['comments'],self.store.detail('xhs','n1')['comments'])
        for bad in (False,{},'xiaohongshu:content:n2','zhihu:content:n1:answer','xiaohongshu:comment:n1:c1',ref+"' OR 1=1"):
            with self.assertRaisesRegex(ValueError,'invalid_results_entity_ref'):self.store.detail('xhs','n1',entity_ref=bad)
            with self.assertRaisesRegex(ValueError,'invalid_results_entity_ref'):self.store.evaluate('xhs','n1',4,entity_ref=bad)
        self.assertEqual(self.store.list_results()['results_revision'],revision)
        with self.assertRaisesRegex(ValueError,'result_not_found'):self.store.detail('xhs','n1',entity_ref='xiaohongshu:dynamic:n1')

    def test_legacy_typed_selector_is_readonly_and_scores_exact_historical_version(self):
        from unittest.mock import patch
        answer={'kind':'content','platform':'zhihu','content_id':'123','content_type':'answer','body':'answer'}
        with patch('results_store._typed_entity',side_effect=lambda entity,_:entity):
            self.store.ingest({'run_id':'legacy-a','records':[answer]})
            self.store.ingest({'run_id':'legacy-b','records':[{**answer,'content_type':'article','body':'article'}]})
        with sqlite3.connect(str(self.store.path)) as db:before=list(db.execute('SELECT * FROM versions'))
        ref='zhihu:content:123:article';detail=self.store.detail('zhihu','123',entity_ref=ref)
        self.assertEqual(len(detail['versions']),1);self.assertEqual(detail['versions'][0]['record']['body'],'article')
        result=self.store.evaluate('zhihu','123',4,entity_ref=ref)
        self.assertEqual(result['version_hash'],detail['contents'][0]['version_hash'])
        self.assertEqual(self.store.detail('zhihu','123',entity_ref='zhihu:content:123:answer')['contents'][0]['owner_evaluations'],[])
        with sqlite3.connect(str(self.store.path)) as db:self.assertEqual(list(db.execute('SELECT * FROM versions')),before)

    def test_dashboard_results_passes_optional_exact_selector(self):
        from types import SimpleNamespace
        import dashboard
        self.store.ingest({'records':[{'kind':'content','platform':'zhihu','content_id':'123','content_type':kind,'body':kind} for kind in ('answer','article')]})
        app=SimpleNamespace(lock=threading.Lock(),results_store=self.store)
        payload={'platform':'zhihu','content_id':'123','entity_ref':'zhihu:content:123:answer'}
        self.assertEqual(dashboard.Application.results(app,'detail',payload)['contents'][0]['body'],'answer')
        self.assertEqual(dashboard.Application.results(app,'evaluate',{**payload,'value':4})['entity_ref'],payload['entity_ref'])

    def test_legacy_zhihu_types_project_without_rewriting_records_hashes_or_evaluations(self):
        from unittest.mock import patch
        data={'run_id':'legacy','records':[{'kind':'content','platform':'zhihu','content_id':'123','content_type':'answer','body':'original answer'}]}
        with patch('results_store._typed_entity',side_effect=lambda entity,_:entity):
            first=self.store.ingest(data)
        initial=self.store.detail('zhihu','123')['contents'][0]
        # Simulate the original stored evaluation entity without rewriting it during compatibility reads.
        with sqlite3.connect(str(self.store.path)) as db:
            db.execute('INSERT INTO evaluations(entity,version_hash,revision,evaluator_ref,value,tags,created_at) VALUES (?,?,?,?,?,?,?)',
                       ('zhihu:content:123',initial['version_hash'],first['results_revision'],'legacy-owner',4,'[]','2026-10-11T00:00:00Z'))
            before=[tuple(row) for row in db.execute('SELECT * FROM versions ORDER BY entity,version_hash')]
            evaluations_before=[tuple(row) for row in db.execute('SELECT * FROM evaluations')]
        with patch('results_store._typed_entity',side_effect=lambda entity,_:entity):
            self.store.ingest({'run_id':'legacy-article','records':[{'kind':'content','platform':'zhihu','content_id':'123','content_type':'article','body':'article'}]})
        detail=self.store.detail('zhihu','123')
        self.assertEqual({r['content_type'] for r in detail['contents']},{'answer','article'})
        answer=next(r for r in detail['contents'] if r['content_type']=='answer')
        self.assertEqual(answer['version_hash'],initial['version_hash']);self.assertEqual(answer['owner_evaluations'][0]['value'],4)
        self.assertEqual(self.store.saved_version('zhihu','123','content',initial['version_hash'])['body'],'original answer')
        # New typed import of the same historical version must not make version lookup ambiguous.
        self.store.ingest({**data,'run_id':'new-typed'})
        self.assertEqual(self.store.saved_version('zhihu','123','content',initial['version_hash'])['first_version_revision'],first['results_revision'])
        current=next(r for r in self.store.detail('zhihu','123')['contents'] if r['content_type']=='answer')
        self.assertEqual(current['owner_evaluations'][0]['value'],4)
        self.assertEqual(current['first_version_revision'],first['results_revision'])
        self.assertEqual(self.store.detail('zhihu','123',first['results_revision'])['stats']['contents'],1)
        with sqlite3.connect(str(self.store.path)) as db:
            old_rows=[tuple(row) for row in db.execute('SELECT * FROM versions WHERE entity=? ORDER BY entity,version_hash',('zhihu:content:123',))]
            self.assertIn(before[0],old_rows)
            self.assertEqual([tuple(row) for row in db.execute('SELECT * FROM evaluations')],evaluations_before)

    def test_metrics_are_observations_not_body_versions(self):
        first = self.store.ingest(imported())
        second = self.store.ingest(imported('run-2', likes=9))
        self.assertEqual(second['added_versions'], 0)
        detail = self.store.detail('xhs', 'n1')
        self.assertEqual(len(detail['versions']), 1)
        self.assertEqual(len(detail['observations']), 2)
        self.assertEqual(detail['contents'][0]['metrics']['liked_count']['reported_value'], 9)
        before = self.store.detail('xhs', 'n1', first['results_revision'])
        self.assertEqual(before['contents'][0]['metrics']['liked_count']['reported_value'], 1)

    def test_content_changes_and_revert_preserve_versions(self):
        self.store.ingest(imported())
        self.store.ingest(imported('run-2', body='Changed'))
        self.store.ingest(imported('run-3'))
        result = self.store.detail('xhs', 'n1')
        self.assertEqual(len(result['versions']), 2)
        self.assertEqual(result['contents'][0]['body'], 'Original text')
        self.assertEqual(len({v['body_hash'] for v in result['versions']}), 2)

    def test_exact_platform_relationships_and_media_manifest(self):
        data = imported()
        data['records'].append({'kind': 'content', 'platform': 'weibo', 'content_id': 'n1', 'body': 'Elsewhere'})
        data['media'][0]['platform'] = 'xiaohongshu'
        self.store.ingest(data)
        result = self.store.detail('xhs', 'n1')
        self.assertEqual(result['stats']['comments'], 2)
        self.assertEqual(result['stats']['root_comments'], 1)
        self.assertEqual(result['stats']['orphan_replies'], 0)
        self.assertEqual(result['authors'][0]['display_name'], 'Author')
        self.assertEqual(len(result['media']), 1)
        self.assertFalse(result['media'][0]['reward_eligible'])

    def test_unmapped_media_is_not_guessed_as_attached(self):
        data = imported()
        del data['media'][0]['content_id']
        self.store.ingest(data)
        self.assertEqual(self.store.detail('xhs', 'n1')['media'], [])
        media = self.store.list_results({'kind': 'media'})['items'][0]
        self.assertIn('content_id', media['missing_fields'])

    def test_missing_parent_and_orphan_not_counted_as_roots(self):
        data = imported()
        data['records'][1]['parent_comment_id'] = None
        data['records'][2]['parent_comment_id'] = 'unknown-comment'
        self.store.ingest(data)
        stats = self.store.detail('xhs', 'n1')['stats']
        self.assertEqual((stats['root_comments'], stats['unknown_parent_comments'], stats['orphan_replies']), (0, 1, 1))

    def test_reported_weibo_self_root_preserves_parent_and_history(self):
        rows = [{'kind':'content','platform':'weibo','content_id':'wb1','body':'body'},
                {'kind':'comment','platform':'weibo','content_id':'wb1','comment_id':'r1',
                 'parent_comment_id':'r1','parent_status':'reported','text':'root'},
                {'kind':'comment','platform':'weibo','content_id':'wb1','comment_id':'c1',
                 'parent_comment_id':'r1','parent_status':'reported','text':'reply'}]
        first = self.store.ingest({'records':rows})['results_revision']
        rows[1]['parent_status'] = 'reported_root'
        self.store.ingest({'records':rows})
        detail = self.store.detail('wb','wb1')
        self.assertEqual((detail['stats']['root_comments'],detail['stats']['orphan_replies']), (1,0))
        self.assertEqual(next(r for r in detail['comments'] if r['comment_id']=='r1')['parent_comment_id'], 'r1')
        self.assertEqual(self.store.detail('wb','wb1',first)['stats']['root_comments'],0)
        # A root label alone must not turn another comment's parent into a root.
        rows[2]['parent_status'] = 'reported_root'
        self.store.ingest({'records':rows})
        self.assertEqual(self.store.detail('wb','wb1')['stats']['root_comments'],1)

    def test_dynamic_supported(self):
        data = imported()
        data['records'][0]['kind'] = 'dynamic'
        self.store.ingest(data)
        self.assertEqual(self.store.detail('xhs', 'n1')['contents'][0]['kind'], 'dynamic')

    def test_owner_evaluation_is_version_bound_and_not_reward(self):
        before = self.store.ingest(imported())['results_revision']
        evaluation = self.store.evaluate('xhs', 'n1', 4, ['useful', 'food'])
        self.assertFalse(evaluation['reward_eligible'])
        self.assertEqual(self.store.list_results({'tags': ['useful']})['total'], 1)
        self.assertEqual(self.store.list_results({'tags': ['useful']}, revision=before)['total'], 0)
        self.store.ingest(imported('run-2', body='New text'))
        self.assertEqual(self.store.detail('xhs', 'n1')['contents'][0]['owner_evaluations'], [])
        with sqlite3.connect(self.store.path) as db:
            names = {row[0] for row in db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        self.assertFalse({'rewards', 'ratings', 'participants'} & names)

    def test_business_tags_use_latest_evaluation_without_erasing_history(self):
        self.store.ingest(imported())
        previous = self.store.evaluate('xhs', 'n1', 2, ['old'])['results_revision']
        self.store.evaluate('xhs', 'n1', 4, ['new'])
        self.assertEqual(self.store.list_results({'tags': ['old']})['total'], 0)
        self.assertEqual(self.store.list_results({'tags': ['old']}, revision=previous)['total'], 1)
        self.assertEqual(len(self.store.detail('xhs', 'n1')['contents'][0]['owner_evaluations']), 2)

    def test_filter_by_author_platform_and_reported_time(self):
        self.store.ingest(imported())
        result = self.store.list_results({'platform': 'xhs', 'author_id': 'a1', 'kind': 'content',
             'published_after': '2024-01-01T00:00:00+00:00', 'published_before': '2024-01-01T08:00:00+08:00'})
        self.assertEqual(result['total'], 1)
        self.assertEqual(result['items'][0]['published_at_utc'], '2024-01-01T00:00:00Z')
        self.assertEqual(self.store.list_results({'published_after': '2025-01-01T00:00:00Z'})['total'], 0)

    def test_unknown_source_time_not_fabricated(self):
        data = imported()
        data['records'][0]['published_at_reported'] = '2024-01-01 12:00:00'
        self.store.ingest(data)
        row = self.store.detail('xhs', 'n1')['contents'][0]
        self.assertIsNone(row['published_at_utc'])
        self.assertIn('published_at_utc', row['missing_fields'])
        self.assertEqual(self.store.list_results({'published_after': '2023-01-01T00:00:00Z'})['total'], 0)

    def test_cursor_freezes_revision_after_concurrent_import(self):
        self.store.ingest(imported())
        first = self.store.list_results(limit=2)
        data = imported('run-2')
        data['records'].append({'kind': 'content', 'platform': 'xiaohongshu', 'content_id': 'n2', 'body': 'Another'})
        self.store.ingest(data)
        second = self.store.list_results(limit=2, cursor=first['next_cursor'])
        self.assertEqual(first['results_revision'], second['results_revision'])
        self.assertEqual(second['total'], 5)
        self.assertFalse({v['entity_ref'] for v in first['items']} & {v['entity_ref'] for v in second['items']})
        with self.assertRaisesRegex(ValueError, 'invalid_results_cursor'):
            self.store.list_results({'kind': 'comment'}, cursor=first['next_cursor'])

    def test_json_jsonl_csv_fixed_revision_and_formula_escape(self):
        data = imported(body=' =HYPERLINK("https://example.com")')
        data['records'][0]['title'] = '\t+SUM(1,2)'
        revision = self.store.ingest(data)['results_revision']
        exported = self.store.export('json', revision=revision)
        self.assertEqual(json.loads(exported['data'])['row_count'], 5)
        self.assertEqual(json.loads(self.store.export('jsonl')['data'].splitlines()[0])['kind'], 'export_manifest')
        rows = list(csv.DictReader(io.StringIO(self.store.export('csv')['data'].decode('utf-8-sig'))))
        content = next(row for row in rows if row['kind'] == 'content')
        self.assertTrue(content['title'].startswith("'\t+"))
        self.assertTrue(content['body'].startswith("' ="))
        self.store.ingest(imported('run-2', body='Updated'))
        self.assertEqual(exported['data'], self.store.export('json', revision=revision)['data'])
        self.assertIn('missing_field_counts', self.store.export('csv')['manifest'])

    def test_sensitive_input_rejected_without_partial_commit(self):
        data = imported()
        data['records'].append({'kind': 'content', 'platform': 'xiaohongshu', 'content_id': 'n2', 'body': 'Cookie=private-value'})
        with self.assertRaisesRegex(ValueError, 'results_sensitive'):
            self.store.ingest(data)
        self.assertEqual(self.store.list_results()['total'], 0)
        data = imported()
        data['records'][0]['author'] = {'access_token': 'fake'}
        with self.assertRaisesRegex(ValueError, 'results_sensitive_field'):
            self.store.ingest(data)

    def test_malformed_text_and_anonymous_media_rejected(self):
        data = imported()
        data['records'][0]['body'] = 123
        with self.assertRaisesRegex(ValueError, 'invalid_results_text_field'):
            self.store.ingest(data)
        data = imported()
        data['media'] = [{'bytes': 2}]
        with self.assertRaisesRegex(ValueError, 'missing_media_identity'):
            self.store.ingest(data)

    def test_validation_errors_are_typed(self):
        self.store.ingest(imported())
        calls = [lambda: self.store.list_results({'unknown': 1}),
                 lambda: self.store.list_results({'published_after': '2024-01-01'}),
                 lambda: self.store.list_results({'tags': 'food'}),
                 lambda: self.store.list_results(limit=True),
                 lambda: self.store.list_results(offset=1),
                 lambda: self.store.list_results(cursor='%%%'),
                 lambda: self.store.list_results(revision=True),
                 lambda: self.store.evaluate('xhs', 'n1', True),
                 lambda: self.store.evaluate('xhs', 'n1', 6),
                 lambda: self.store.evaluate('xhs', 'n1', 4, tags=''),
                 lambda: self.store.export('html')]
        for call in calls:
            with self.subTest(call=call), self.assertRaises(ValueError):
                call()

    def test_parallel_import_is_single_logical_commit(self):
        results, errors = [], []
        def run():
            try:
                results.append(ResultsStore(self.temporary.name).ingest(imported()))
            except Exception as exc:
                errors.append(exc)
        threads = [threading.Thread(target=run) for _ in range(4)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join()
        self.assertEqual(errors, [])
        self.assertEqual(sum(not row['duplicate_import'] for row in results), 1)
        self.assertEqual(self.store.list_results()['results_revision'], 1)

    def test_symlink_database_rejected(self):
        another = Path(self.temporary.name) / 'other'
        another.mkdir()
        results = another / 'results'
        results.mkdir()
        (results / 'results.sqlite3').symlink_to(self.store.path)
        with self.assertRaisesRegex(ValueError, 'unsafe_results_database'):
            ResultsStore(another)

    def test_xlsx_when_installed_or_explicit_missing_dependency(self):
        self.store.ingest(imported(body='\ufeff=SUM(1,2)'))
        try:
            result = self.store.export('xlsx')
        except ValueError as exc:
            self.assertEqual(str(exc), 'xlsx_dependency_missing')
            return
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(result['data']))
        self.assertEqual(workbook.sheetnames, ['results', 'manifest'])
        self.assertEqual(workbook['results'].max_row, 6)
        header = [cell.value for cell in workbook['results'][1]]
        body_column = header.index('body')
        content_row = next(row for row in workbook['results'].iter_rows(min_row=2) if row[0].value == 'content')
        self.assertTrue(content_row[body_column].value.startswith("'\ufeff="))
        self.assertEqual(content_row[body_column].data_type, 's')


if __name__ == '__main__':
    unittest.main()
