"""Local versioned business results. No source requests, cloud ACKs or rewards."""
import base64
from contextlib import closing
import csv
from datetime import datetime, timezone
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import re
import sqlite3

try:
    from .mc_runner import PLATFORMS, contains_secret
except ImportError:
    from mc_runner import PLATFORMS, contains_secret


SCHEMA = 'crowd-local-results-v1'
KINDS = ('content', 'dynamic', 'comment', 'author', 'media')
META_KEYS = {'metrics', 'origin', 'source', 'source_kind', 'coverage', 'record_hash',
             'reward_eligible', 'upstream_transport', 'collected_at', 'observed_at'}
SECRET_KEYS = re.compile(r'(?i)^(cookie|authorization|access_token|refresh_token|xsec_token|token|password|passwd|secret|sessionid|api_key|private_key)$')
FILTER_KEYS = {'platform', 'author_id', 'content_id', 'published_after', 'published_before', 'tags', 'kind'}
CSV_FIELDS = ('kind', 'platform', 'content_id', 'comment_id', 'parent_comment_id',
              'author_id', 'title', 'body', 'text', 'published_at_reported',
              'published_at_utc', 'version_hash', 'body_hash', 'metrics', 'source',
              'coverage', 'missing_fields', 'owner_evaluations', 'record')


def _json(value):
    try:
        return json.dumps(value, ensure_ascii=False, sort_keys=True,
                          separators=(',', ':'), allow_nan=False)
    except (ValueError, TypeError, UnicodeError):
        raise ValueError('invalid_results_encoding') from None


def _sha(value):
    return hashlib.sha256(_json(value).encode('utf8')).hexdigest()


def _identifier(value, name='identifier'):
    if isinstance(value, int) and not isinstance(value, bool):
        value = str(value)
    if not isinstance(value, str) or not re.fullmatch(r'[A-Za-z0-9_.:-]{1,200}', value) or contains_secret(value):
        raise ValueError('invalid_' + name)
    return value


def _platform(value):
    if isinstance(value, str) and value in PLATFORMS:
        return PLATFORMS[value]
    if isinstance(value, str) and value in (*PLATFORMS.values(), 'reddit'):
        return value
    raise ValueError('invalid_platform')


def _safe(value, depth=0):
    if depth > 12:
        raise ValueError('results_nesting_limit')
    if isinstance(value, dict):
        if len(value) > 300:
            raise ValueError('results_object_limit')
        for key, item in value.items():
            if not isinstance(key, str) or len(key) > 200 or SECRET_KEYS.fullmatch(key):
                raise ValueError('results_sensitive_field')
            _safe(item, depth + 1)
    elif isinstance(value, list):
        if len(value) > 10000:
            raise ValueError('results_array_limit')
        for item in value:
            _safe(item, depth + 1)
    elif isinstance(value, str):
        if len(value) > 100000 or contains_secret(value):
            raise ValueError('results_sensitive_or_oversized_text')
    elif value is not None and not isinstance(value, (int, float, bool)):
        raise ValueError('invalid_results_value')


def _date(value, strict=False):
    try:
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            seconds = value / 1000 if abs(value) > 100000000000 else value
            parsed = datetime.fromtimestamp(seconds, timezone.utc)
        elif isinstance(value, str):
            parsed = datetime.fromisoformat(value.replace('Z', '+00:00'))
            if parsed.tzinfo is None:
                raise ValueError()
            parsed = parsed.astimezone(timezone.utc)
        else:
            raise ValueError()
        return parsed.isoformat(timespec='seconds').replace('+00:00', 'Z')
    except (ValueError, OverflowError, OSError):
        if strict:
            raise ValueError('invalid_published_filter') from None
        return None


def _tags(value):
    if not isinstance(value, list) or len(value) > 20:
        raise ValueError('invalid_tags')
    if any(not isinstance(v, str) or not v.strip() or len(v) > 80 or contains_secret(v) for v in value):
        raise ValueError('invalid_tags')
    return sorted(set(value))


def _filters(filters):
    if filters is None:
        return {}
    if not isinstance(filters, dict) or set(filters) - FILTER_KEYS:
        raise ValueError('invalid_results_filters')
    out = dict(filters)
    if 'platform' in out:
        out['platform'] = _platform(out['platform'])
    for field in ('author_id', 'content_id'):
        if field in out:
            out[field] = _identifier(out[field], field)
    for field in ('published_after', 'published_before'):
        if field in out:
            if not isinstance(out[field], str):
                raise ValueError('invalid_published_filter')
            out[field] = _date(out[field], strict=True)
    if out.get('published_after', '') > out.get('published_before', '9999'):
        raise ValueError('invalid_published_range')
    if 'tags' in out:
        out['tags'] = _tags(out['tags'])
    if 'kind' in out and out['kind'] not in KINDS:
        raise ValueError('invalid_results_kind')
    return out


def _record(record, default_platform=None):
    if not isinstance(record, dict) or record.get('kind') not in KINDS:
        raise ValueError('invalid_results_record')
    _safe(record)
    item = dict(record)
    item['platform'] = _platform(item.get('platform', default_platform))
    kind = item['kind']
    for field in ('title', 'body', 'text', 'content_type', 'display_name'):
        if item.get(field) is not None and not isinstance(item[field], str):
            raise ValueError('invalid_results_text_field')
    author = item.get('author') if isinstance(item.get('author'), dict) else {}
    aid = item.get('author_id') or item.get('creator_id') or item.get('user_id') or author.get('author_id') or author.get('user_id') or author.get('id')
    aid = _identifier(aid, 'author_id') if aid is not None else None
    cid = _identifier(item['content_id'], 'content_id') if item.get('content_id') is not None else None
    if kind in ('content', 'dynamic', 'comment') and cid is None:
        raise ValueError('missing_content_id')
    if kind == 'comment':
        suffix = _identifier(item.get('comment_id'), 'comment_id')
        item['comment_id'] = suffix
        if item.get('parent_comment_id') not in (None, ''):
            item['parent_comment_id'] = _identifier(item['parent_comment_id'], 'parent_comment_id')
        identity = item['platform'] + ':comment:' + cid + ':' + suffix
    elif kind == 'author':
        if aid is None:
            raise ValueError('missing_author_id')
        identity = item['platform'] + ':author:' + aid
    elif kind == 'media':
        mid = item.get('media_id') or item.get('sha256')
        if not mid and (not isinstance(item.get('path'), str) or not item['path']):
            raise ValueError('missing_media_identity')
        mid = _identifier(mid, 'media_id') if mid else _sha({'path': item.get('path'), 'content_id': cid})
        identity = item['platform'] + ':media:' + mid
    else:
        identity = item['platform'] + ':' + kind + ':' + cid
    item['content_id'], item['author_id'] = cid, aid
    reported = item.get('published_at_reported', item.get('published_at'))
    item['published_at_utc'] = _date(reported)
    expected = ('body', 'author_id', 'published_at_utc') if kind in ('content', 'dynamic') else ('text', 'parent_comment_id') if kind == 'comment' else ('content_id',) if kind == 'media' else ('display_name',)
    missing = item.get('missing_fields', [])
    if not isinstance(missing, list) or len(missing) > 100 or any(not isinstance(key, str) or len(key) > 100 for key in missing):
        raise ValueError('invalid_missing_fields')
    item['missing_fields'] = sorted(set(missing) | {key for key in expected if item.get(key) is None})
    semantic = {key: value for key, value in item.items() if key not in META_KEYS}
    version_hash = _sha(semantic)
    body_hash = _sha({key: item.get(key) for key in ('title', 'body', 'text', 'content_type')})
    return identity, item, semantic, version_hash, body_hash


class ResultsStore:
    """A private local SQLite store, deliberately separate from cloud/reward tables."""
    def __init__(self, root):
        base = Path(root).expanduser()
        if base.is_symlink():
            raise ValueError('unsafe_results_root')
        base.mkdir(parents=True, exist_ok=True, mode=0o700)
        self.root = base.resolve() / 'results'
        if self.root.is_symlink():
            raise ValueError('unsafe_results_root')
        self.root.mkdir(mode=0o700, exist_ok=True)
        self.path = self.root / 'results.sqlite3'
        if self.path.is_symlink():
            raise ValueError('unsafe_results_database')
        fd = os.open(self.path, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600) if not self.path.exists() else None
        if fd is not None:
            os.close(fd)
        os.chmod(self.path, 0o600)
        with closing(self._connect()) as db:
            db.executescript('''
                CREATE TABLE IF NOT EXISTS state (id INTEGER PRIMARY KEY CHECK(id=1), revision INTEGER NOT NULL);
                INSERT OR IGNORE INTO state VALUES (1,0);
                CREATE TABLE IF NOT EXISTS imports (digest TEXT PRIMARY KEY, revision INTEGER NOT NULL, run_id TEXT);
                CREATE TABLE IF NOT EXISTS versions (
                    entity TEXT NOT NULL, version_hash TEXT NOT NULL, kind TEXT NOT NULL, platform TEXT NOT NULL,
                    content_id TEXT, author_id TEXT, published_at TEXT, body_hash TEXT NOT NULL,
                    record TEXT NOT NULL, revision INTEGER NOT NULL,
                    PRIMARY KEY(entity,version_hash));
                CREATE INDEX IF NOT EXISTS version_lookup ON versions(platform,kind,content_id,revision);
                CREATE INDEX IF NOT EXISTS version_author ON versions(platform,author_id,published_at,revision);
                CREATE TABLE IF NOT EXISTS observations (
                    entity TEXT NOT NULL, version_hash TEXT NOT NULL, import_digest TEXT NOT NULL,
                    revision INTEGER NOT NULL, observed_at TEXT NOT NULL, record TEXT NOT NULL,
                    PRIMARY KEY(entity,version_hash,import_digest));
                CREATE INDEX IF NOT EXISTS observation_revision ON observations(entity,revision);
                CREATE TABLE IF NOT EXISTS evaluations (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, entity TEXT NOT NULL, version_hash TEXT NOT NULL,
                    revision INTEGER NOT NULL, evaluator_ref TEXT NOT NULL, value INTEGER NOT NULL,
                    tags TEXT NOT NULL, created_at TEXT NOT NULL);
                CREATE INDEX IF NOT EXISTS evaluation_version ON evaluations(entity,version_hash,revision);
            ''')

    def _connect(self):
        if self.path.is_symlink():
            raise ValueError('unsafe_results_database')
        db = sqlite3.connect(str(self.path), timeout=10)
        db.row_factory = sqlite3.Row
        db.execute('PRAGMA busy_timeout=10000')
        db.execute('PRAGMA journal_mode=DELETE')
        return db

    def _revision(self, db, revision):
        latest = db.execute('SELECT revision FROM state WHERE id=1').fetchone()[0]
        if revision is None:
            return latest
        if type(revision) is not int or revision < 0 or revision > latest:
            raise ValueError('invalid_results_revision')
        return revision

    def ingest(self, imported):
        if not isinstance(imported, dict) or not isinstance(imported.get('records'), list):
            raise ValueError('invalid_imported_results')
        records, media = imported['records'], imported.get('media', [])
        if not isinstance(media, list) or len(records) > 10000 or len(media) > 1000:
            raise ValueError('results_input_limit')
        # The input is already a business projection. Never import browser/session configuration.
        run_id = imported.get('run_id')
        if run_id is not None:
            run_id = _identifier(run_id, 'run_id')
        source = imported.get('source', {})
        _safe(source)
        platforms = {_platform(row.get('platform')) for row in records if isinstance(row, dict)}
        default_platform = next(iter(platforms)) if len(platforms) == 1 else None
        prepared = [_record(row) for row in records]
        prepared += [_record({**row, 'kind': 'media'}, default_platform) for row in media if isinstance(row, dict)]
        if len(prepared) != len(records) + len(media):
            raise ValueError('invalid_media_record')
        fingerprint = _sha({'run_id': run_id, 'source': source, 'records': [row[1] for row in prepared]})
        if len(_json([row[1] for row in prepared]).encode('utf8')) > 20 * 1024 * 1024:
            raise ValueError('results_input_byte_limit')
        now = datetime.now(timezone.utc).isoformat(timespec='seconds').replace('+00:00', 'Z')
        with closing(self._connect()) as db, db:
            db.execute('BEGIN IMMEDIATE')
            duplicate = db.execute('SELECT revision FROM imports WHERE digest=?', (fingerprint,)).fetchone()
            if duplicate:
                return {'results_revision': self._revision(db, None), 'import_revision': duplicate[0],
                        'added_versions': 0, 'added_observations': 0, 'duplicate_import': True}
            revision = self._revision(db, None) + 1
            versions = observations = 0
            for entity, item, semantic, version_hash, body_hash in prepared:
                result = db.execute('INSERT OR IGNORE INTO versions VALUES (?,?,?,?,?,?,?,?,?,?)',
                    (entity, version_hash, item['kind'], item['platform'], item['content_id'], item['author_id'],
                     item['published_at_utc'], body_hash, _json(semantic), revision))
                versions += result.rowcount
                observation = {'metrics': item.get('metrics', {}), 'origin': item.get('origin'),
                    'source': source, 'run_id': run_id, 'source_kind': item.get('source_kind', imported.get('source_kind')),
                    'coverage': item.get('coverage', imported.get('coverage', 'unverified')),
                    'upstream_transport': item.get('upstream_transport'), 'reward_eligible': False}
                observation_digest = _sha({'import': fingerprint, 'observation': observation})
                result = db.execute('INSERT OR IGNORE INTO observations VALUES (?,?,?,?,?,?)',
                    (entity, version_hash, observation_digest, revision, _date(item.get('observed_at')) or now, _json(observation)))
                observations += result.rowcount
            db.execute('INSERT INTO imports VALUES (?,?,?)', (fingerprint, revision, run_id))
            db.execute('UPDATE state SET revision=? WHERE id=1', (revision,))
        return {'results_revision': revision, 'added_versions': versions,
                'added_observations': observations, 'duplicate_import': False}

    def _rows(self, db, filters, revision):
        # Latest observation selects the active version, including a return to an older text.
        sql = '''SELECT v.*,o.record AS observation,o.observed_at FROM observations o
            JOIN versions v ON v.entity=o.entity AND v.version_hash=o.version_hash
            WHERE o.revision<=? AND NOT EXISTS
            (SELECT 1 FROM observations n WHERE n.entity=o.entity AND n.revision<=?
             AND (n.revision>o.revision OR (n.revision=o.revision AND n.rowid>o.rowid)))'''
        values = [revision, revision]
        for field, column in (('platform', 'platform'), ('author_id', 'author_id'), ('content_id', 'content_id'), ('kind', 'kind')):
            if field in filters:
                sql += ' AND v.' + column + '=?'
                values.append(filters[field])
        for field, operator in (('published_after', '>='), ('published_before', '<=')):
            if field in filters:
                sql += ' AND v.published_at' + operator + '?'
                values.append(filters[field])
        rows = []
        for row in db.execute(sql + ' ORDER BY v.entity', values):
            evaluations = [dict(e) for e in db.execute('SELECT * FROM evaluations WHERE entity=? AND version_hash=? AND revision<=? ORDER BY revision,id', (row['entity'], row['version_hash'], revision))]
            for evaluation in evaluations:
                evaluation['tags'] = json.loads(evaluation['tags'])
                evaluation['type'] = 'quality'
                evaluation['reward_eligible'] = False
            current_by_evaluator = {e['evaluator_ref']: e for e in evaluations}
            if filters.get('tags') and not set(filters['tags']).issubset({tag for e in current_by_evaluator.values() for tag in e['tags']}):
                continue
            item = {**json.loads(row['record']), **json.loads(row['observation']),
                    'entity_ref': row['entity'], 'version_hash': row['version_hash'],
                    'body_hash': row['body_hash'], 'first_version_revision': row['revision'],
                    'observed_at': row['observed_at'], 'owner_evaluations': evaluations}
            rows.append(item)
            if len(rows) > 50000:
                raise ValueError('results_export_row_limit')
        return rows

    def list_results(self, filters=None, revision=None, limit=50, offset=0, cursor=None):
        filters = _filters(filters)
        if type(limit) is not int or not 1 <= limit <= 200 or type(offset) is not int or offset < 0 or offset > 50000:
            raise ValueError('invalid_results_page')
        if cursor is not None:
            try:
                if not isinstance(cursor, str) or len(cursor) > 2000:
                    raise ValueError()
                state = json.loads(base64.urlsafe_b64decode(cursor.encode('ascii')))
                if set(state) != {'revision', 'offset', 'filters_hash'} or state['filters_hash'] != _sha(filters):
                    raise ValueError()
                if revision is not None and revision != state['revision']:
                    raise ValueError()
                revision, offset = state['revision'], state['offset']
                if type(offset) is not int or not 0 <= offset <= 50000:
                    raise ValueError()
            except (ValueError, TypeError, UnicodeError, KeyError):
                raise ValueError('invalid_results_cursor') from None
        elif offset and revision is None:
            raise ValueError('results_revision_required_for_offset')
        with closing(self._connect()) as db:
            revision = self._revision(db, revision)
            items = self._rows(db, filters, revision)
        end = offset + limit
        next_cursor = base64.urlsafe_b64encode(_json({'revision': revision, 'offset': end, 'filters_hash': _sha(filters)}).encode()).decode() if end < len(items) else None
        return {'schema': SCHEMA, 'results_revision': revision, 'items': items[offset:end],
                'total': len(items), 'next_cursor': next_cursor, 'coverage': 'observed_local_imports_only'}

    def saved_version(self, platform, content_id, kind, version_hash):
        platform, content_id = _platform(platform), _identifier(content_id, 'content_id')
        if kind not in KINDS or not isinstance(version_hash, str) or not re.fullmatch(r'[a-f0-9]{64}', version_hash):
            raise ValueError('invalid_results_version')
        with closing(self._connect()) as db:
            versions = db.execute('SELECT revision FROM versions WHERE platform=? AND content_id=? AND kind=? AND version_hash=?',
                                  (platform, content_id, kind, version_hash)).fetchall()
            if len(versions) != 1:
                raise ValueError('results_version_not_found')
            records = self._rows(db, {'platform': platform, 'content_id': content_id, 'kind': kind}, versions[0]['revision'])
        matching = [row for row in records if row['version_hash'] == version_hash]
        if len(matching) != 1:
            raise ValueError('results_version_not_found')
        return matching[0]

    def detail(self, platform, content_id, revision=None):
        platform, content_id = _platform(platform), _identifier(content_id, 'content_id')
        with closing(self._connect()) as db:
            revision = self._revision(db, revision)
            rows = self._rows(db, {'platform': platform, 'content_id': content_id}, revision)
            contents = [row for row in rows if row['kind'] in ('content', 'dynamic')]
            if not contents:
                raise ValueError('result_not_found')
            author_ids = {row['author_id'] for row in contents if row.get('author_id')}
            authors = [row for row in self._rows(db, {'platform': platform, 'kind': 'author'}, revision) if row.get('author_id') in author_ids]
            versions = [dict(row) for row in db.execute('SELECT version_hash,body_hash,revision,record FROM versions WHERE platform=? AND content_id=? AND kind IN (\'content\',\'dynamic\') AND revision<=? ORDER BY revision', (platform, content_id, revision))]
            observations = [dict(row) for row in db.execute('SELECT o.revision,o.observed_at,o.record,o.version_hash FROM observations o JOIN versions v ON v.entity=o.entity AND v.version_hash=o.version_hash WHERE v.platform=? AND v.content_id=? AND v.kind IN (\'content\',\'dynamic\') AND o.revision<=? ORDER BY o.revision,o.rowid', (platform, content_id, revision))]
        for row in versions + observations:
            row['record'] = json.loads(row['record'])
        comments = [row for row in rows if row['kind'] == 'comment']
        roots = sum(row.get('parent_comment_id') in ('0', '') for row in comments)
        known = {row['comment_id'] for row in comments}
        orphans = sum(row.get('parent_comment_id') not in (None, '', '0') and row['parent_comment_id'] not in known for row in comments)
        return {'schema': SCHEMA, 'results_revision': revision, 'contents': contents, 'comments': comments,
                'authors': authors, 'media': [row for row in rows if row['kind'] == 'media'],
                'versions': versions, 'observations': observations,
                'stats': {'contents': len(contents), 'comments': len(comments), 'root_comments': roots,
                          'unknown_parent_comments': sum(row.get('parent_comment_id') is None for row in comments),
                          'orphan_replies': orphans, 'coverage': 'observed_local_imports_only'}}

    def evaluate(self, platform, content_id, value, tags=None, evaluator_ref='local-owner', revision=None):
        platform, content_id = _platform(platform), _identifier(content_id, 'content_id')
        if type(value) is not int or not 0 <= value <= 5:
            raise ValueError('invalid_owner_evaluation')
        evaluator_ref = _identifier(evaluator_ref, 'evaluator_ref')
        tags = _tags([] if tags is None else tags)
        now = datetime.now(timezone.utc).isoformat(timespec='seconds').replace('+00:00', 'Z')
        with closing(self._connect()) as db, db:
            db.execute('BEGIN IMMEDIATE')
            selected = self._revision(db, revision)
            rows = self._rows(db, {'platform': platform, 'content_id': content_id}, selected)
            contents = [row for row in rows if row['kind'] in ('content', 'dynamic')]
            if len(contents) != 1:
                raise ValueError('content_evaluation_requires_one_content')
            content = contents[0]
            new_revision = self._revision(db, None) + 1
            row = db.execute('INSERT INTO evaluations(entity,version_hash,revision,evaluator_ref,value,tags,created_at) VALUES (?,?,?,?,?,?,?)', (content['entity_ref'], content['version_hash'], new_revision, evaluator_ref, value, _json(tags), now))
            db.execute('UPDATE state SET revision=? WHERE id=1', (new_revision,))
        return {'evaluation_id': row.lastrowid, 'results_revision': new_revision,
                'version_hash': content['version_hash'], 'value': value, 'tags': tags,
                'reward_eligible': False, 'type': 'quality'}

    def export(self, format='json', filters=None, revision=None):
        if format not in ('json', 'jsonl', 'csv', 'xlsx'):
            raise ValueError('invalid_export_format')
        filters = _filters(filters)
        with closing(self._connect()) as db:
            revision = self._revision(db, revision)
            items = self._rows(db, filters, revision)
        manifest = {'schema': SCHEMA, 'results_revision': revision, 'row_count': len(items),
                    'filters': filters, 'coverage': 'observed_local_imports_only',
                    'missing_field_counts': {field: sum(field in row['missing_fields'] for row in items)
                                             for field in sorted({field for row in items for field in row['missing_fields']})},
                    'columns': list(CSV_FIELDS), 'formula_escape': 'leading_apostrophe_for_formula_like_text'}
        if format == 'json':
            data = (_json({**manifest, 'items': items}) + '\n').encode('utf8')
            mime = 'application/json; charset=utf-8'
        elif format == 'jsonl':
            data = ('\n'.join([_json({'kind': 'export_manifest', **manifest})] + [_json(row) for row in items]) + '\n').encode('utf8')
            mime = 'application/x-ndjson; charset=utf-8'
        else:
            flat = []
            for item in items:
                row = {field: item.get(field) for field in CSV_FIELDS}
                row['record'] = item
                values = []
                for field in CSV_FIELDS:
                    value = row[field]
                    text = _json(value) if isinstance(value, (dict, list)) else '' if value is None else str(value)
                    if text.lstrip(' \t\r\n\v\f\ufeff').startswith(('=', '+', '-', '@')):
                        text = "'" + text
                    values.append(text)
                flat.append(values)
            if format == 'csv':
                output = io.StringIO(newline='')
                writer = csv.writer(output)
                writer.writerow(CSV_FIELDS)
                writer.writerows(flat)
                data, mime = output.getvalue().encode('utf-8-sig'), 'text/csv; charset=utf-8'
            else:
                if importlib.util.find_spec('openpyxl') is None:
                    raise ValueError('xlsx_dependency_missing')
                from openpyxl import Workbook
                from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
                if any(len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value) for row in flat for value in row):
                    raise ValueError('xlsx_cell_limit_or_invalid_character')
                book = Workbook(write_only=True)
                sheet = book.create_sheet('results')
                sheet.append(CSV_FIELDS)
                for row in flat:
                    sheet.append(row)
                metadata = book.create_sheet('manifest')
                for key, value in manifest.items():
                    metadata.append([key, _json(value)])
                output = io.BytesIO()
                book.save(output)
                data, mime = output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        return {'data': data, 'mime': mime, 'filename': 'results-r' + str(revision) + '.' + format,
                'sha256': hashlib.sha256(data).hexdigest(), 'row_count': len(items),
                'results_revision': revision, 'schema': SCHEMA, 'manifest': manifest}
